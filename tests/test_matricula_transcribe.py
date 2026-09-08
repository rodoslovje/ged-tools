import json
import os
import tempfile

import openpyxl
import pytest

from tools import matricula_to_json as mj
from tools import matricula_transcribe as mt

BOOK_URL = "https://data.matricula-online.eu/sl/slovenia/ljubljana/podzemelj/04455/?pg=1"

MARRIAGE_ROW = {
    "marriage_date": "1874-04-19",
    "address": "Gradac 1",
    "groom_name": "Henrik",
    "groom_surname": "Prelesnik",
    "groom_alt_surname": "Prelesnik",
    "bride_address": "Gradac 51",
    "bride_name": "Marija",
    "bride_surname": "Horošec",
    "bride_alt_surname": "",
    "notes": "ženin iz ž. Semič",
    "uncertain": ["groom_surname"],
}

BIRTH_ROW = {
    "birth_date": "1856-02-28",
    "baptism_date": "1856-02-28",
    "address": "Krašnji vrh 15",
    "child_name": "Marko",
    "father_name": "Marko",
    "father_surname": "Šavor",
    "father_alt_surname": "Šaver",
    "mother_name": "Katarina",
    "mother_surname": "Stankovič",
    "mother_alt_surname": "",
    "notes": "umrl 03.12.1914",
    "uncertain": [],
}


# --- URL and book metadata ------------------------------------------------

def test_normalize_book_url_pins_page_one():
    assert mt.normalize_book_url(
        "https://data.matricula-online.eu/sl/slovenia/ljubljana/podzemelj/04455/?pg=7"
    ) == BOOK_URL
    assert mt.normalize_book_url(
        "https://data.matricula-online.eu/sl/slovenia/ljubljana/podzemelj/04455"
    ) == BOOK_URL


def test_normalize_book_url_rejects_other_hosts():
    with pytest.raises(SystemExit):
        mt.normalize_book_url("https://example.com/sl/slovenia/x/1/")


def test_page_url_rewrites_the_page_number():
    assert mt.page_url(BOOK_URL, 42).endswith("/04455/?pg=42")


def test_parse_book_page_decodes_the_image_list():
    html = (
        "<title>Poročna knjiga / Trauungsbuch - 04455 | Podzemelj</title>"
        '<script>new arc.imageview.MatriculaDocView("document", {'
        '"labels": ["1", "2"], '
        # base64 of '.../P Podzemelj 1873-1922/P Podzemelj 1873-1922-0001.jpg'
        '"files": ["/image/'
        "aHR0cDovL2gvUCUyMFBvZHplbWVsaiUyMDE4NzMtMTkyMi9QJTIwUG9kemVtZWxqJTIw"
        "MTg3My0xOTIyLTAwMDEuanBn"
        '/"], "path": "https://img.data.matricula-online.eu"});</script>'
    )
    title, images = mt.parse_book_page(html)
    assert "Poročna knjiga" in title
    assert images == [
        "http://h/P%20Podzemelj%201873-1922/P%20Podzemelj%201873-1922-0001.jpg"
    ]
    assert mt.book_stem(images[0]) == "P Podzemelj 1873-1922"


def test_parse_book_page_without_an_image_list():
    with pytest.raises(SystemExit):
        mt.parse_book_page("<title>Podzemelj</title><p>no viewer here</p>")


@pytest.mark.parametrize("stem,expected", [
    # The archive names birth folders R (rojstna); the index sheets use K.
    ("R Podzemelj 1675-1703", ("K", "Podzemelj", "1675-1703")),
    ("P Podzemelj 1873-1922", ("P", "Podzemelj", "1873-1922")),
    ("M Podzemelj 1725-1768", ("M", "Podzemelj", "1725-1768")),
    ("M Podzemelj 1728-1803-nemški viteški red", ("M", "Podzemelj", "1728-1803")),
    ("K Stari trg pri Ložu 1659-1677", ("K", "Stari trg pri Ložu", "1659-1677")),
    ("nekaj drugega", (None, "", "")),
])
def test_parse_stem(stem, expected):
    assert mt.parse_stem(stem) == expected


@pytest.mark.parametrize("title,path,expected", [
    ("Krstna knjiga / Taufbuch - 01723 | Podzemelj", "", "K"),
    ("Poročna knjiga / Trauungsbuch - 04455 | Podzemelj", "", "P"),
    ("Mrliška knjiga / Sterbebuch - 01732 | Podzemelj", "", "M"),
    # No marker in the title: fall back to the scan folder's leading letter.
    ("04104 | Podzemelj", "http://h/R%20Podzemelj%201871-1885/x-0001.jpg", "K"),
    ("04104 | Podzemelj", "", None),
])
def test_detect_kind(title, path, expected):
    assert mt.detect_kind(title, path) == expected


def test_parish_from_url():
    assert mt.parish_from_url(BOOK_URL) == "Podzemelj"
    assert mt.parish_from_url(
        "https://data.matricula-online.eu/sl/slovenia/ljubljana/stari-trg/04455/"
    ) == "Stari Trg"


# --- page selection -------------------------------------------------------

@pytest.mark.parametrize("spec,expected", [
    ("", [1, 2, 3, 4, 5]),
    ("2-4", [2, 3, 4]),
    ("1,3,5", [1, 3, 5]),
    ("1-2,2-3", [1, 2, 3]),
    ("4-99", [4, 5]),
    (" 3 ", [3]),
])
def test_parse_pages(spec, expected):
    assert mt.parse_pages(spec, 5) == expected


@pytest.mark.parametrize("spec", ["abc", "4-2", "99-100"])
def test_parse_pages_rejects_bad_input(spec):
    with pytest.raises(SystemExit):
        mt.parse_pages(spec, 5)


# --- row assembly ---------------------------------------------------------

def test_build_rows_fills_in_what_the_model_is_not_asked_for():
    rows = mt.build_rows("P", [(6, {"rows": [MARRIAGE_ROW]})],
                         "Podzemelj", BOOK_URL, "Renko_Luka", True)
    fields = [field for field, _header in mt.MARRIAGE_COLUMNS]
    row = dict(zip(fields, rows[0]))
    assert row["parish"] == "Podzemelj"
    assert row["url"] == mt.page_url(BOOK_URL, 6)
    assert row["interpret"] == "Renko_Luka"
    # The row's position on its page, which the model is not asked for.
    assert row["seq"] == 1


def test_uncertain_fields_are_named_in_their_own_column():
    rows = mt.build_rows("P", [(6, {"rows": [MARRIAGE_ROW]})],
                         "Podzemelj", BOOK_URL, "Renko_Luka", True)
    fields = [f for f, _h in mt.columns_for("P", True)]
    row = dict(zip(fields, rows[0]))
    # The note keeps only what the register said; the flag stands apart, so
    # clearing it when the index is ready cannot take a real note with it.
    assert row["notes"] == "ženin iz ž. Semič"
    assert row["flags"] == "priimek ženina"


def test_no_flags_leaves_the_column_out_entirely():
    rows = mt.build_rows("P", [(6, {"rows": [MARRIAGE_ROW]})],
                         "Podzemelj", BOOK_URL, "Renko_Luka", False)
    assert len(rows[0]) == len(mt.MARRIAGE_COLUMNS)
    fields = [f for f, _h in mt.MARRIAGE_COLUMNS]
    assert "flags" not in fields
    assert rows[0][fields.index("notes")] == "ženin iz ž. Semič"


def test_a_flag_without_a_note_leaves_the_note_empty():
    row = dict(MARRIAGE_ROW, notes="", uncertain=["marriage_date"])
    got = bride("P", row)
    assert got["notes"] is None
    assert got["flags"] == "datum poroke"


def test_uncertain_field_left_empty_is_not_flagged():
    row = dict(MARRIAGE_ROW, notes="", uncertain=["bride_alt_surname"])
    got = bride("P", row)
    assert got["flags"] is None


def test_alt_surname_repeating_the_modern_one_is_dropped():
    # MARRIAGE_ROW's groom alt surname equals his surname.
    rows = mt.build_rows("P", [(6, {"rows": [MARRIAGE_ROW]})],
                         "Podzemelj", BOOK_URL, "X", True)
    fields = [f for f, _h in mt.MARRIAGE_COLUMNS]
    assert rows[0][fields.index("groom_alt_surname")] is None


def test_a_real_alt_surname_is_kept():
    rows = mt.build_rows("K", [(2, {"rows": [BIRTH_ROW]})],
                         "Radovica", BOOK_URL, "X", True)
    fields = [f for f, _h in mt.BIRTH_COLUMNS]
    assert rows[0][fields.index("father_alt_surname")] == "Šaver"


def test_pages_without_entries_contribute_nothing():
    payload = {"page_note": "prazna stran", "rows": []}
    assert mt.build_rows("P", [(1, payload)], "Podzemelj", BOOK_URL, "X", True) == []


def test_rows_keep_page_order_and_seq_restarts_on_each_page():
    # This is the convention the hand-made indexes follow: 'zp. št.' is the
    # entry's position on its page, so it starts again at 1 on the next one.
    rows = mt.build_rows(
        "P",
        [(6, {"rows": [MARRIAGE_ROW, MARRIAGE_ROW, MARRIAGE_ROW]}),
         (7, {"rows": [MARRIAGE_ROW, MARRIAGE_ROW]})],
        "Podzemelj", BOOK_URL, "X", True,
    )
    fields = [f for f, _h in mt.MARRIAGE_COLUMNS]
    seqs = [row[fields.index("seq")] for row in rows]
    pages = [row[fields.index("url")][-1] for row in rows]
    assert seqs == [1, 2, 3, 1, 2]
    assert pages == ["6", "6", "6", "7", "7"]


def test_seq_is_not_asked_of_the_model():
    for kind in ("K", "P"):
        assert "seq" not in mt.MODEL_FIELDS[kind]


# --- schema ---------------------------------------------------------------

@pytest.mark.parametrize("kind", ["K", "P"])
def test_schema_asks_for_exactly_the_model_owned_columns(kind):
    schema = mt.build_schema(kind)
    row = schema["properties"]["rows"]["items"]
    sheet_fields = {field for field, _header in mt.COLUMNS[kind]}
    asked = set(row["properties"]) - {"uncertain"}
    assert asked == set(mt.MODEL_FIELDS[kind])
    # Everything the model returns must have a home in the sheet.
    assert asked <= sheet_fields
    # The columns this script fills in itself are not asked of the model.
    assert sheet_fields - asked == {"seq", "parish", "url", "interpret"}
    assert row["required"] == mt.MODEL_FIELDS[kind] + ["uncertain"]
    assert row["additionalProperties"] is False


# --- the workbook matricula_to_json has to be able to read ----------------

@pytest.mark.parametrize("kind,row,expected_kind", [
    ("P", MARRIAGE_ROW, "P"),
    ("K", BIRTH_ROW, "K"),
])
def test_written_workbook_round_trips_through_matricula_to_json(
    kind, row, expected_kind
):
    rows = mt.build_rows(kind, [(6, {"rows": [row]})],
                         "Podzemelj", BOOK_URL, "Renko_Luka", False)
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, f"Indeks {kind} Podzemelj - 1873-1922.xlsx")
        mt.write_workbook(path, kind, rows, f"{kind} 1873-1922")

        assert mj.detect_book_kind(path) == expected_kind
        read = list(mj.read_rows(path))
        assert len(read) == 1
        sheet_kind, record = read[0]
        assert sheet_kind == expected_kind
        assert record["parish"] == "Podzemelj"
        assert record["url"] == mt.page_url(BOOK_URL, 6)


def test_marriage_workbook_produces_the_expected_family_record():
    rows = mt.build_rows("P", [(6, {"rows": [MARRIAGE_ROW]})],
                         "Podzemelj", BOOK_URL, "Renko_Luka", False)
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "Indeks P Podzemelj - 1873-1922.xlsx")
        mt.write_workbook(path, "P", rows, "P")
        _kind, record = next(iter(mj.read_rows(path)))

    family = mj.marriage_record(record)
    assert family["husband"]["name"] == "Henrik"
    assert family["husband"]["surname"] == "Prelesnik"
    assert "alt_surname" not in family["husband"]
    assert family["wife"]["surname"] == "Horošec"
    assert family["marriage"]["date"] == "19 APR 1874"
    assert family["marriage"]["place"] == "Podzemelj, Gradac 1"
    assert family["links"] == [mt.page_url(BOOK_URL, 6)]


def test_birth_workbook_produces_the_expected_person_record():
    rows = mt.build_rows("K", [(2, {"rows": [BIRTH_ROW]})],
                         "Radovica", BOOK_URL, "Renko_Luka", False)
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "Indeks K Radovica - 1856-1900.xlsx")
        mt.write_workbook(path, "K", rows, "K")
        _kind, record = next(iter(mj.read_rows(path)))

    person = mj.birth_record(record)
    assert person["name"] == "Marko"
    assert person["surname"] == "Šavor"
    assert person["birth"]["date"] == "28 FEB 1856"
    assert person["birth"]["place"] == "Radovica, Krašnji vrh 15"
    # The death date in the notes is picked up by matricula_to_json.
    assert person["death"]["date"] == "3 DEC 1914"
    father, mother = person["parents_list"]
    assert father["alt_surname"] == "Šaver"
    assert "alt_surname" not in mother


def test_workbook_header_matches_the_declared_columns():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "Indeks P X - 1800-1900.xlsx")
        mt.write_workbook(path, "P", [], "P")
        sheet = openpyxl.load_workbook(path).worksheets[0]
        header = [cell.value for cell in sheet[1]]
    assert header == [h for _f, h in mt.MARRIAGE_COLUMNS]


def test_the_flag_column_is_appended_after_the_template_columns():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "Indeks P X - 1800-1900.xlsx")
        mt.write_workbook(path, "P", [], "P", add_flags=True)
        sheet = openpyxl.load_workbook(path).worksheets[0]
        header = [cell.value for cell in sheet[1]]
    # The template columns keep their order and position; 'preveri' goes last,
    # so deleting it leaves exactly the workbook a human indexer would write.
    assert header == [h for _f, h in mt.MARRIAGE_COLUMNS] + ["preveri"]


def test_matricula_to_json_ignores_the_flag_column():
    rows = mt.build_rows("P", [(6, {"rows": [MARRIAGE_ROW]})],
                         "Podzemelj", BOOK_URL, "Renko_Luka", True)
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "Indeks P Podzemelj - 1873-1922.xlsx")
        mt.write_workbook(path, "P", rows, "P", add_flags=True)
        header = list(openpyxl.load_workbook(path).worksheets[0][1])
        assert mj._resolve_columns([c.value for c in header])[-1] == ""
        _kind, record = next(iter(mj.read_rows(path)))
    assert "preveri" not in json.dumps(record, ensure_ascii=False).lower()
    assert record["notes"] == "ženin iz ž. Semič"


def test_safe_filename_keeps_the_readable_name():
    assert mt.safe_filename("Indeks P Stari trg/Lož - 1659-1677") == \
        "Indeks P Stari trg-Lož - 1659-1677"


# --- run-stopping API errors ----------------------------------------------

def test_workspace_error_explains_how_to_pass_the_workspace():
    text = mt.explain_api_error(RuntimeError(
        "Error code: 400 - anthropic-workspace-id is required when "
        "authenticating with an identity-linked API key"
    ))
    assert "ANTHROPIC_WORKSPACE_ID" in text
    assert "--workspace" in text


def test_credit_error_points_at_billing():
    text = mt.explain_api_error(RuntimeError("Your credit balance is too low"))
    assert "credit" in text.lower()


def test_unrecognised_error_still_says_what_to_do():
    text = mt.explain_api_error(RuntimeError("something else went wrong"))
    assert "something else went wrong" in text
    assert "cached pages are not re-billed" in text


# --- alt surnames the model made up rather than read -----------------------

@pytest.mark.parametrize("surname,alt", [
    # ć is not a Slovenian letter and 19th-century clerks did not write it;
    # an alt that differs from the surname only by one is an invention.
    ("Jakofčič", "Jakofčić"),
    ("Šopčič", "Šopčić"),
    ("Žunič", "Žunić"),
    ("Jakofčič", "Jakofčič"),
])
def test_invented_alt_surname_is_dropped(surname, alt):
    row = dict(MARRIAGE_ROW, groom_surname=surname, groom_alt_surname=alt)
    rows = mt.build_rows("P", [(6, {"rows": [row]})],
                         "Podzemelj", BOOK_URL, "X", False)
    fields = [f for f, _h in mt.MARRIAGE_COLUMNS]
    assert rows[0][fields.index("groom_alt_surname")] is None


@pytest.mark.parametrize("surname,alt", [
    # A real reading of the book's own spelling has to survive.
    ("Pašič", "Pasič"),
    ("Žunič", "Zunič"),
    ("Kovačič", "Kowatschitsch"),
    ("Šavor", "Šaver"),
    ("Grdešič", "Gerdešič"),
])
def test_genuine_alt_surname_is_kept(surname, alt):
    row = dict(MARRIAGE_ROW, groom_surname=surname, groom_alt_surname=alt)
    rows = mt.build_rows("P", [(6, {"rows": [row]})],
                         "Podzemelj", BOOK_URL, "X", False)
    fields = [f for f, _h in mt.MARRIAGE_COLUMNS]
    assert rows[0][fields.index("groom_alt_surname")] == alt


def test_the_prompt_names_the_letters_that_do_not_belong():
    assert "ć" in mt.SYSTEM_PROMPT
    assert "č, š and ž" in mt.SYSTEM_PROMPT


@pytest.mark.parametrize("german,slovene", [
    ("Möttling", "Metlika"),
    ("Tschernembl", "Črnomelj"),
    ("Gottschee", "Kočevje"),
    ("Laibach", "Ljubljana"),
    ("Klagenfurt", "Celovec"),
    ("Cilli", "Celje"),
])
def test_prompt_maps_the_austrian_place_names_to_slovenian(german, slovene):
    # The registers were kept under Austrian administration and name places in
    # German; the indexes are written with today's Slovenian names.
    assert f"{german} -> {slovene}" in mt.SYSTEM_PROMPT


def test_prompt_keeps_the_registers_own_form_in_brackets():
    # The house style the other contributors follow, e.g. 'Celje (Cilli) 53'.
    assert "'Celje (Cilli) 53'" in mt.SYSTEM_PROMPT


# --- what the notes are told to leave out ---------------------------------

def test_notes_rules_drop_the_defaults_and_the_priest():
    prompt = mt.SYSTEM_PROMPT
    # 'samski'/'zakonski' are what the columns already say; the priest
    # identifies nobody. Only the departures from the ordinary are recorded.
    assert "never write 'samski' / 'samska'" in prompt
    assert "never write 'zakonski'" in prompt
    assert "Never record the priest" in prompt
    assert "vdovec" in prompt and "nezakonski" in prompt


def test_marriage_notes_are_ordered_groom_then_bride():
    prompt = mt.KIND_PROMPT["P"]
    assert "groom first, then bride" in prompt
    groom = prompt.index("Ženin:")
    bride = prompt.index("Nevesta:")
    witnesses = prompt.index("Priči:")
    assert groom < bride < witnesses


def test_birth_notes_ask_for_godparents():
    assert "Botra" in mt.KIND_PROMPT["K"]


# --- widow brides ---------------------------------------------------------
#
# A register enters a remarrying widow under the surname she carries that day
# — her dead husband's. The index files her under her maiden name, which the
# same entry gives through her parents, and keeps the married one in the alt
# column.

def widow_row(**over):
    row = dict(MARRIAGE_ROW, bride_name="Ana", bride_surname="Filak",
               bride_alt_surname="", uncertain=[],
               notes="Ženin: 27 let; starša Jožef Zemelj in Ana Kralj. "
                     "Nevesta: vdova, 35 let, kajžarka, r. Vardijan; starša "
                     "Marko Vardijan, kmet, in Rozalija Vanjaš.")
    row.update(over)
    return row


def bride(kind, row):
    rows = mt.build_rows(kind, [(6, {"rows": [row]})],
                         "Podzemelj", BOOK_URL, "X", True)
    fields = [f for f, _h in mt.columns_for(kind, True)]
    return dict(zip(fields, rows[0]))


def test_widow_bride_is_filed_under_her_maiden_surname():
    got = bride("P", widow_row())
    assert got["bride_surname"] == "Vardijan"      # her father's
    assert got["bride_alt_surname"] == "Filak"     # what the register wrote


def test_the_swap_is_flagged_for_a_human():
    got = bride("P", widow_row())
    assert "priimek neveste" in got["flags"]
    assert "preveri" not in (got["notes"] or "")


def test_widow_already_under_her_maiden_name_is_left_alone():
    row = widow_row(bride_surname="Vardijan", bride_alt_surname="Filak")
    got = bride("P", row)
    assert got["bride_surname"] == "Vardijan"
    assert got["bride_alt_surname"] == "Filak"
    assert "priimek neveste" not in (got["flags"] or "")


def test_a_bride_who_is_not_a_widow_is_untouched():
    row = widow_row(notes="Nevesta: 22 let; starša Marko Vardijan, kmet, in "
                          "Rozalija Vanjaš.")
    assert bride("P", row)["bride_surname"] == "Filak"


def test_a_widowed_mother_does_not_make_the_bride_a_widow():
    # 'Barbara vdova Šterk roj. Plut' describes the bride's mother.
    row = widow_row(notes="Nevesta: 26 let, kovača hči; starša † Franc Repovž, "
                          "kovač, in Barbara vdova Šterk roj. Plut.")
    assert bride("P", row)["bride_surname"] == "Filak"


def test_a_widowed_witness_does_not_make_the_bride_a_widow():
    row = widow_row(notes="Nevesta: 34 let; starša Martin Žunič in Ana Bilič. "
                          "Priči: Marija Skubic, meščanka, vdova.")
    assert bride("P", row)["bride_surname"] == "Filak"


def test_widow_noted_before_the_groom_segment_is_still_caught():
    # Some entries open with 'Nevesta vdova ...' ahead of 'Ženin:'.
    row = widow_row(
        bride_surname="Klepec",
        notes="Nevesta vdova (por. Straus), posestnica. Ženin: 23 let; starša "
              "Janez Starašinič, posestnik, in Ana roj. Zugelj. Nevesta: 31 "
              "let; starša Nikolaj Klepec, posestnik, in † Barbara roj. Tomc.")
    got = bride("P", row)
    assert got["bride_surname"] == "Klepec"   # father's; already correct here
    assert got["bride_alt_surname"] == "Straus"   # the husband she survived


def test_a_displaced_book_spelling_is_kept_in_the_notes():
    row = widow_row(bride_surname="Štrucelj", bride_alt_surname="Strucelj",
                    notes="Nevesta: vdova, 68 let; starša Matija Klobučar, "
                          "polovični kmet, in Ana Horvat.")
    got = bride("P", row)
    assert got["bride_surname"] == "Klobučar"
    assert got["bride_alt_surname"] == "Štrucelj"
    assert "v knjigi Strucelj" in got["notes"]


def test_no_parents_means_no_maiden_name_and_no_change():
    row = widow_row(notes="Nevesta: vdova.")
    assert bride("P", row)["bride_surname"] == "Filak"


def test_the_rule_does_not_touch_birth_books():
    row = dict(BIRTH_ROW, mother_surname="Filak",
               notes="Nevesta: vdova; starša Marko Vardijan, kmet, in Ana Kralj.")
    assert bride("K", row)["mother_surname"] == "Filak"


def test_the_prompt_states_the_widow_rule():
    prompt = mt.KIND_PROMPT["P"]
    assert "MAIDEN surname" in prompt
    assert "AT THIS WEDDING" in prompt
    assert "Vardijan" in prompt and "Filak" in prompt


def test_married_surname_is_put_in_the_alt_column():
    row = widow_row(bride_surname="Karasinič", bride_alt_surname="",
                    notes="Nevesta: vdova po Strucelju, 37 let; starša Jurij "
                          "Karasinič, užitkar, in Katarina Kure.")
    got = bride("P", row)
    assert got["bride_surname"] == "Karasinič"
    assert got["bride_alt_surname"] == "Strucelj"    # locative undone


@pytest.mark.parametrize("locative,expected,vocab", [
    ("Požeku", "Požek", {}),
    ("Marentiču", "Marentič", {}),
    ("Strucelju", "Strucelj", {}),
    ("Kukarju", "Kukar", {}),
    ("Stieblerju", "Stiebler", {}),
    ("Jakši", "Jakša", {}),
    # '-cu' is ambiguous; the book's own spelling of its families decides.
    ("Tomcu", "Tomc", {"Tomc": 181, "Tomec": 12}),
    ("Kostelcu", "Kostelec", {"Kostelec": 2}),
    ("Milavcu", "Milavec", {"Milavec": 12}),
])
def test_locative_is_turned_back_into_a_plain_surname(locative, expected, vocab):
    assert mt.nominative(locative, vocab) == expected


def test_an_invented_alt_does_not_block_the_married_surname():
    # The model sometimes fills the alt slot with a ć variant it made up; that
    # must be cleared before the widow's married surname claims the column.
    row = widow_row(bride_surname="Karasinič", bride_alt_surname="Karasinić",
                    notes="Nevesta: vdova po Strucelju, 37 let; starša Jurij "
                          "Karasinič, užitkar, in Katarina Kure.")
    assert bride("P", row)["bride_alt_surname"] == "Strucelj"


def test_a_real_spelling_gives_way_to_the_married_surname_but_is_kept():
    row = widow_row(bride_surname="Anžlovar", bride_alt_surname="Anzlover",
                    notes="Nevesta: vdova, 32 let, vdova po Gorše; starša Anton "
                          "Anžlovar, kmet, in Jera Valentin.")
    got = bride("P", row)
    assert got["bride_alt_surname"] == "Gorše"
    assert "v knjigi Anzlover" in got["notes"]


def test_the_registers_own_nee_annotation_is_understood():
    # 'Kralj p. Gorše' — married name, then her maiden one. Her father is a
    # Gorše, so 'Kralj' is the husband's.
    row = widow_row(bride_surname="Gorše", bride_alt_surname="Kralj p. Gorše",
                    notes="Nevesta: vdova, 38 let, hišna posestnica; starša "
                          "Andrej Gorše, dninar, in Terezija Anžlovar.")
    got = bride("P", row)
    assert got["bride_surname"] == "Gorše"
    assert got["bride_alt_surname"] == "Kralj"


def test_widow_clause_before_the_groom_segment_survives_an_abbreviation():
    # '(por. Straus)' — the clause must not be cut at the dot in 'por.'
    row = widow_row(
        bride_surname="Klepec", bride_alt_surname="",
        notes="Nevesta vdova (por. Straus), posestnica. Ženin: 23 let; starša "
              "Janez Starašinič, posestnik, in Ana roj. Zugelj. Nevesta: 31 "
              "let; starša Nikolaj Klepec, posestnik, in † Barbara roj. Tomc.")
    assert bride("P", row)["bride_alt_surname"] == "Straus"


def test_surname_vocabulary_counts_names_from_columns_and_notes():
    vocab = mt.surname_vocabulary([
        {"groom_surname": "Tomc", "bride_surname": "Tomc",
         "notes": "starša Nikolaj Tomc, polhubnik, in Barbara Saje."},
    ])
    assert vocab["Tomc"] == 3
    assert vocab["Saje"] == 1
